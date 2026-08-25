"""Rebuilds the exercise catalog workbook the artwork skill drives from.

The workbook lives next to the illustrations (default: the "Slike vaj" folder on
the Desktop) and is the ordered list the artwork generator walks. Regenerate it
whenever assets/data/exercises.json changes, so exercises added or removed there
show up in the artwork queue.

    python tool/build_exercise_catalog_xlsx.py [--out PATH] [--images DIR]

An "Ima sliko" column marks which slugs already ship an illustration, so the
remaining work is visible in the sheet itself.
"""

from __future__ import annotations

import argparse
import collections
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
CATALOG = REPO_ROOT / "assets" / "data" / "exercises.json"
DEFAULT_DIR = Path.home() / "OneDrive" / "Desktop" / "Slike vaj"
IMAGE_SUFFIXES = (".png", ".jpg", ".jpeg", ".webp")

HEADERS = [
    "Zap. št.",
    "Ime vaje",
    "Slug",
    "Alias",
    "Telesni del",
    "Kategorija",
    "Vzorec gibanja",
    "Modalnost",
    "Oprema",
    "Primarna mišica",
    "Sekundarne mišice",
    "Stabilizatorji",
    "CNS ocena",
    "Vpliv na regeneracijo",
    "Metrika beleženja",
    "Podpira obteženo telesno težo",
    "Privzeti počitek (s)",
    "Izpeljana različica",
    "Ima sliko",
]
WIDTHS = [8, 34, 32, 26, 22, 16, 18, 14, 18, 20, 30, 24, 11, 20, 18, 28, 18, 18, 11]

TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F6F78")


def join(value) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value)
    return str(value)


def yes_no(value) -> str:
    return "Da" if value else "Ne"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path, default=DEFAULT_DIR / "herculex-vaje.xlsx")
    parser.add_argument("--images", type=Path, default=DEFAULT_DIR)
    args = parser.parse_args()

    catalog = json.loads(CATALOG.read_text(encoding="utf-8"))
    have_image = set()
    if args.images.is_dir():
        have_image = {
            p.stem.lower()
            for p in args.images.iterdir()
            if p.is_file() and p.suffix.lower() in IMAGE_SUFFIXES
        }

    wb = Workbook()
    ws = wb.active
    ws.title = "Vse vaje"

    missing = [e for e in catalog if e["slug"].lower() not in have_image]
    ws.cell(row=1, column=1, value="Herculex — katalog vaj").font = TITLE_FONT
    ws.cell(
        row=2,
        column=1,
        value=(
            f"Vse vaje in različice iz assets/data/exercises.json. "
            f"Skupaj: {len(catalog)} · s sliko: {len(catalog) - len(missing)} · "
            f"brez slike: {len(missing)}"
        ),
    )

    for col, (header, width) in enumerate(zip(HEADERS, WIDTHS), start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    for index, e in enumerate(catalog, start=1):
        ws.append(
            [
                index,
                e.get("name"),
                e.get("slug"),
                join(e.get("aka")),
                e.get("bodyPart"),
                e.get("category"),
                e.get("movementPatternRaw") or e.get("movementPattern"),
                e.get("modality"),
                e.get("equipment"),
                e.get("primaryMuscle"),
                join(e.get("secondaryMuscles")),
                join(e.get("stabilizers")),
                e.get("cnsScore"),
                e.get("recoveryImpact"),
                e.get("loggingMetric"),
                yes_no(e.get("supportsWeightedBodyweight")),
                e.get("defaultRestSeconds"),
                yes_no(e.get("derived")),
                yes_no(e["slug"].lower() in have_image),
            ]
        )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(HEADERS))}{4 + len(catalog)}"

    summary = wb.create_sheet("Povzetek")
    summary.cell(row=1, column=1, value="Herculex — povzetek kataloga").font = TITLE_FONT
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 14

    rows: list[tuple[str, object]] = [
        ("Meritev", "Vrednost"),
        ("Skupno število vaj / različic", len(catalog)),
        ("S sliko", len(catalog) - len(missing)),
        ("Brez slike", len(missing)),
        ("Izpeljane različice", sum(1 for e in catalog if e.get("derived"))),
        ("Osnovni zapisi", sum(1 for e in catalog if not e.get("derived"))),
    ]
    for label, key in (
        ("Po telesnem delu", "bodyPart"),
        ("Po kategoriji", "category"),
        ("Po opremi", "equipment"),
    ):
        counts = collections.Counter(e.get(key) or "—" for e in catalog)
        rows += [("", ""), ("", ""), (label, "Število")]
        rows += sorted(counts.items())

    for offset, (label, value) in enumerate(rows, start=3):
        summary.cell(row=offset, column=1, value=label or None)
        summary.cell(row=offset, column=2, value=value if value != "" else None)
        if value in ("Vrednost", "Število"):
            summary.cell(row=offset, column=1).font = Font(bold=True)
            summary.cell(row=offset, column=2).font = Font(bold=True)

    missing_sheet = wb.create_sheet("Brez slike")
    missing_sheet.column_dimensions["A"].width = 40
    missing_sheet.column_dimensions["B"].width = 40
    missing_sheet.column_dimensions["C"].width = 20
    for col, header in enumerate(("Ime vaje", "Slug", "Oprema"), start=1):
        cell = missing_sheet.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for e in missing:
        missing_sheet.append([e.get("name"), e.get("slug"), e.get("equipment")])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"wrote {args.out} — {len(catalog)} rows, {len(missing)} without artwork")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
